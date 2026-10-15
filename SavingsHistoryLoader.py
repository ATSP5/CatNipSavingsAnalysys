from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from AnalysisSettings import AnalysisSettings
from InputFileLocator import InputFileLocator
from SavingsDeposit import SavingsDeposit
from SavingsHistory import SavingsHistory


class SavingsHistoryLoader:
    """Czyta historię wpłat (kolumny Data, Kwota) z pierwszego arkusza pliku Excel."""

    def __init__(self, locator: InputFileLocator, settings: AnalysisSettings):
        self.locator = locator
        self.file_name = settings.history_file

    def history(self) -> SavingsHistory:
        try:
            path = self.locator.find_input(self.file_name)
        except FileNotFoundError as error:
            print("[Historia] %s — kapitał startowy 0 zł." % error, flush=True)
            return SavingsHistory(deposits=(), source="brak pliku")
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(c).strip().lower() if c is not None else "" for c in next(rows)]
        day_col, amount_col = header.index("data"), header.index("kwota")
        deposits = []
        for row in rows:
            if row is None or row[day_col] is None or row[amount_col] is None:
                continue
            deposits.append(SavingsDeposit(day=self._to_date(row[day_col]), amount=float(row[amount_col])))
        deposits.sort(key=lambda d: d.day)
        history = SavingsHistory(deposits=tuple(deposits), source=path.name)
        print("[Historia] %d wpłat z %s, razem %.2f zł" % (len(deposits), path.name, history.total), flush=True)
        return history

    @staticmethod
    def _to_date(value) -> date:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)):            # numer seryjny Excela
            return date(1899, 12, 30) + timedelta(days=int(value))
        return datetime.strptime(str(value).strip()[:10], "%Y-%m-%d").date()
